"""Rebuild corridors/east-arapahoe.js from the raw exports in "corridor data/".

Inputs (git-ignored), under "corridor data/East Arapahoe Rd/":
    TMC_Identification.csv                      INRIX TMC geometry + lengths
    East-Arapahoe-Road.csv                      INRIX 5-min speeds per TMC
    2087472_Arapahoe_Road/                      StreetLight Corridor Study export
        2087472_Arapahoe_Road_corridor_study.csv
        Shapefile/2087472_Arapahoe_Road_osm_segment.zip
    ../Statewide Crash Data/*.xlsx              CDOT statewide crash listings

Output (committed):
    corridors/east-arapahoe.js

StreetLight: 160 directional OSM segments (geometry from the export shapefile,
parsed with a minimal stdlib reader) with All-Days ADT / speed / free-flow /
TTI / 85th pct / congested flag plus hourly volume+speed profiles.
INRIX: per-TMC day-average speed + reference (free-flow) speed on straight
start->end chords. Crashes: statewide listings buffer-filtered to the
corridor (see filter_crashes_to_corridor.py; default 250 ft).

Run from anywhere:
    python corridor-studies/scripts/build_east_arapahoe.py [--buffer-ft 250]
"""
import argparse, csv, datetime, glob, json, os, re, struct, zipfile
from collections import defaultdict

import openpyxl

from filter_crashes_to_corridor import filter_crashes

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "..", "corridor data", "East Arapahoe Rd")
CRASH_DIR = os.path.join(HERE, "..", "corridor data", "Statewide Crash Data")
STL = os.path.join(DATA, "2087472_Arapahoe_Road")
OUT = os.path.join(HERE, "..", "corridors", "east-arapahoe.js")

ap = argparse.ArgumentParser()
ap.add_argument("--buffer-ft", type=float, default=250.0)
a = ap.parse_args()


# ── minimal shapefile reader (polylines + dbf), stdlib only ───────────
def read_shapefile_zip(zip_path):
    """Returns [(record_attrs_dict, [[lon,lat], ...]), ...] from a zipped shapefile."""
    zp = zipfile.ZipFile(zip_path)
    shp = zp.read([n for n in zp.namelist() if n.endswith(".shp")][0])
    dbf = zp.read([n for n in zp.namelist() if n.endswith(".dbf")][0])
    # dbf: header -> field descriptors -> fixed-width records
    nrec = struct.unpack("<I", dbf[4:8])[0]
    hsz, rsz = struct.unpack("<HH", dbf[8:12])
    fields, p = [], 32
    while dbf[p] != 0x0D:
        name = dbf[p:p + 11].split(b"\x00")[0].decode()
        fields.append((name, dbf[p + 16]))
        p += 32
    attrs = []
    for i in range(nrec):
        rec = dbf[hsz + i * rsz: hsz + (i + 1) * rsz][1:]  # skip deletion flag
        row, off = {}, 0
        for name, flen in fields:
            row[name] = rec[off:off + flen].decode("latin-1").strip()
            off += flen
        attrs.append(row)
    # shp: 100-byte header, then records (big-endian header, little-endian content)
    lines, p = [], 100
    while p < len(shp):
        clen = struct.unpack(">I", shp[p + 4:p + 8])[0] * 2
        body = shp[p + 8: p + 8 + clen]
        p += 8 + clen
        shape_type = struct.unpack("<i", body[0:4])[0]
        if shape_type in (3, 5, 13, 15):   # PolyLine / Polygon (+Z variants)
            n_parts, n_pts = struct.unpack("<ii", body[36:44])
            pts_off = 44 + 4 * n_parts
            pts = [list(struct.unpack("<dd", body[pts_off + 16 * k: pts_off + 16 * k + 16]))
                   for k in range(n_pts)]
            lines.append([[round(x, 6), round(y, 6)] for x, y in pts])
        else:
            lines.append([])
    return list(zip(attrs, lines))


# ── StreetLight corridor study -> directional segment features ────────
HOURLY_RE = re.compile(r"^\d+: \d+(?:am|pm) \((.+)\)$")   # hourly parts only
stl_geom = {r["segment_id"]: line
            for r, line in read_shapefile_zip(
                os.path.join(STL, "Shapefile", "2087472_Arapahoe_Road_osm_segment.zip"))}

by_zone = defaultdict(dict)   # zone id -> day part -> row
zone_meta = {}
with open(os.path.join(STL, "2087472_Arapahoe_Road_corridor_study.csv"),
          newline="", encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        if not r["Day Type"].startswith("0:"):
            continue
        zid = r["Zone ID"]
        by_zone[zid][r["Day Part"]] = r
        zone_meta[zid] = (r["Zone Cardinal Direction"],
                          float(r["Line Zone Length (Miles)"]))

sl_feats, hourly_acc = [], defaultdict(lambda: defaultdict(lambda: [0.0, 0.0, 0]))
max_adt = 0
for zid, parts in by_zone.items():
    card, length = zone_meta[zid]
    key = "EAST" if card == "EAST" else "WEST"
    short = "EB" if key == "EAST" else "WB"
    allday = parts.get("00: All Day (12am-12am)")
    geom = stl_geom.get(zid)
    if not allday or not geom:
        continue
    adt = round(float(allday["Average Daily Segment Traffic (StL Volume)"]))
    max_adt = max(max_adt, adt)
    hourly = []
    for dp in sorted(p for p in parts if HOURLY_RE.match(p)):
        row = parts[dp]
        label = HOURLY_RE.match(dp).group(1)
        v = round(float(row["Average Daily Segment Traffic (StL Volume)"]))
        s = float(row["Avg Segment Speed (mph)"])
        hourly.append({"h": label, "v": v, "s": s})
        acc = hourly_acc[key][dp]
        acc[0] += v; acc[1] += s * v; acc[2] += 1
    sl_feats.append({
        "type": "Feature",
        "geometry": {"type": "LineString", "coordinates": geom},
        "properties": {
            "name": f"E Arapahoe Rd ({short}) · {length:.2f} mi",
            f"ADT_{key}": adt,
            f"speed_{key}": float(allday["Avg Segment Speed (mph)"]),
            f"freeflow_{key}": float(allday["Free Flow Speed (mph)"]),
            f"tti_{key}": float(allday["Travel Time Index"]),
            f"p85spd_{key}": float(allday["85th Speed Percentile"]),
            f"congested_{key}": allday["Congested Segment"].strip().lower() == "true",
            "length_mi": length,
            "hourly": {key: hourly},
        },
    })

sl_hourly = {}
for key, dps in hourly_acc.items():
    prof = []
    for dp in sorted(dps):
        v_sum, sv_sum, n = dps[dp]
        prof.append({"h": HOURLY_RE.match(dp).group(1),
                     "v": round(v_sum / n),
                     "s": round(sv_sum / v_sum, 1) if v_sum else 0})
    sl_hourly[key] = prof
adt_cap = (max_adt // 5000 + 1) * 5000
print(f"StreetLight: {len(sl_feats)} segments, max ADT {max_adt}, hourly dirs {list(sl_hourly)}")

# ── INRIX: per-TMC day-average + reference speed ──────────────────────
speeds, refs = defaultdict(list), defaultdict(list)
with open(os.path.join(DATA, "East-Arapahoe-Road.csv"), newline="", encoding="utf-8-sig") as f:
    for row in csv.DictReader(f):
        try:
            speeds[row["tmc_code"]].append(float(row["speed"]))
            refs[row["tmc_code"]].append(float(row["reference_speed"]))
        except (KeyError, ValueError):
            continue

feats = []
with open(os.path.join(DATA, "TMC_Identification.csv"), newline="", encoding="utf-8-sig") as f:
    for t in csv.DictReader(f):
        tmc = t["tmc"]
        if tmc not in speeds:
            continue
        key = "EAST" if t["direction"] == "EASTBOUND" else "WEST"
        short = "EB" if key == "EAST" else "WB"
        feats.append({
            "type": "Feature",
            "geometry": {"type": "LineString", "coordinates": [
                [float(t["start_longitude"]), float(t["start_latitude"])],
                [float(t["end_longitude"]),   float(t["end_latitude"])]]},
            "properties": {
                "name": f"Arapahoe Rd @ {t['intersection']} ({short})",
                "tmc": tmc,
                "length_mi": round(float(t["miles"]), 3),
                f"speed_{key}": round(sum(speeds[tmc]) / len(speeds[tmc]), 1),
                f"freeflow_{key}": round(sum(refs[tmc]) / len(refs[tmc]), 1),
            },
        })
print(f"INRIX: {len(feats)} TMC segments")

# ── TCDS short counts: MS2 "Volume Count Report" xlsx exports ─────────
# Station coordinates: Arapahoe stations sit on INRIX TMC on-road points
# just W/O and E/O Parker Rd; Jordan Rd station is ~100 m south of the
# geocoded Arapahoe & Jordan intersection (39.59488, -104.82101).
TCDS_COORDS = {
    "103876": (39.59481, -104.80454, "Arapahoe Rd W/O Parker Rd (SH 83)"),
    "203231": (39.59504, -104.80049, "Arapahoe Rd E/O Parker Rd (SH 83)"),
    "902108": (39.59400, -104.82101, "Jordan Rd S/O Arapahoe Rd"),
}

def parse_tcds_folder(folder):
    stations = {}
    for p in sorted(glob.glob(os.path.join(folder, "*.xlsx"))):
        wb = openpyxl.load_workbook(p, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        kv = {}
        for r in rows:
            cells = [c for c in r if c is not None]
            for i in range(len(cells) - 1):
                k = str(cells[i]).strip()
                if k in ("Location ID", "Direction", "Start Date", "Located On") and k not in kv:
                    kv[k] = cells[i + 1]
        locid = str(kv.get("Location ID", "")).strip()
        base = locid.split("_")[0]
        dirn = str(kv.get("Direction", "") or locid.split("_")[-1]).strip().upper()
        sd = kv.get("Start Date")
        if hasattr(sd, "strftime"):
            date = sd.strftime("%Y-%m-%d")
        else:
            date = datetime.datetime.strptime(str(sd).strip(), "%m/%d/%Y").strftime("%Y-%m-%d")
        vols = [int(r[1]) for r in rows
                if r and isinstance(r[0], str) and re.match(r"^\d{2}:00 - \d{2}:00$", r[0].strip())]
        st = stations.setdefault(base, {"counts": {}})
        st["counts"].setdefault(date, {"hours": list(range(24)), "dirs": {}})["dirs"][dirn] = vols
        st["located"] = str(kv.get("Located On", "")).strip()
    out = []
    for base, st in sorted(stations.items()):
        if base not in TCDS_COORDS:
            print(f"  WARNING: no coordinates mapped for TCDS station {base} ({st.get('located')}) — skipped;"
                  " add it to TCDS_COORDS")
            continue
        lat, lon, label = TCDS_COORDS[base]
        out.append({"id": base, "label": label, "lat": lat, "lon": lon, "counts": st["counts"]})
    return out

tcds_stations = parse_tcds_folder(os.path.join(DATA, "TCDS"))
print(f"TCDS: {len(tcds_stations)} stations: " +
      ", ".join(s["id"] + " (" + "/".join(sorted({d for c in s['counts'].values() for d in c['dirs']})) + ")"
                for s in tcds_stations))

# ── StreetLight O-D (analysis 2087595): zone centroids + compact flows ─
OD_DIR = os.path.join(DATA, "2087595_Arapahoe_Road_OD")

def _abbrev(name):
    s = name.split(" / ")[0]
    for a, b in (("East ", "E "), ("West ", "W "), ("North ", "N "), ("South ", "S "),
                 ("Road", "Rd"), ("Street", "St"), ("Avenue", "Ave"), ("Parkway", "Pkwy"),
                 ("Boulevard", "Blvd"), ("Drive", "Dr"), ("primary_link", "Ramp")):
        s = s.replace(a, b)
    return s

def _compass(deg):
    try: deg = float(deg) % 360
    except (TypeError, ValueError): return ""
    return "NB" if deg < 45 or deg >= 315 else "EB" if deg < 135 else "SB" if deg < 225 else "WB"

def od_zone_list(kind):
    """[(full zone name, label, lat, lon)] from the origin/destination shapefile."""
    zp = os.path.join(OD_DIR, "Shapefile", f"2087595_Arapahoe_Road_OD_{kind}.zip")
    out = []
    for attrs, pts in read_shapefile_zip(zp):
        if not pts:
            continue
        lat = sum(p[1] for p in pts) / len(pts)
        lon = sum(p[0] for p in pts) / len(pts)
        name = attrs["name"]
        wid = name.split(" / ")[-1] if " / " in name else ""
        label = _abbrev(name) + " " + _compass(attrs.get("direction")) + (" ·" + wid[-4:] if wid else "")
        out.append((name, label.strip(), round(lat, 6), round(lon, 6)))
    return out

od_cfg = None
if os.path.isdir(OD_DIR):
    o_zones = od_zone_list("origin")
    d_zones = od_zone_list("destination")
    # dbf names are 30-char truncated; index by prefix for the CSV join
    def zone_index(zones):
        return {name: i for i, (name, _, _, _) in enumerate(zones)}
    def find_zone(idx_map, zones, full_name):
        if full_name in idx_map:
            return idx_map[full_name]
        for i, (name, _, _, _) in enumerate(zones):
            if full_name.startswith(name) or name.startswith(full_name):
                return i
        return None
    o_idx, d_idx = zone_index(o_zones), zone_index(d_zones)
    day_types, day_parts, flows = [], [], []
    missed = set()
    with open(os.path.join(OD_DIR, "2087595_Arapahoe_Road_OD_od_all.csv"),
              newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            oi = find_zone(o_idx, o_zones, r["Origin Zone Name"])
            di = find_zone(d_idx, d_zones, r["Destination Zone Name"])
            if oi is None or di is None:
                missed.add((r["Origin Zone Name"], r["Destination Zone Name"]))
                continue
            dt, dp = r["Day Type"], r["Day Part"]
            if dt not in day_types: day_types.append(dt)
            if dp not in day_parts: day_parts.append(dp)
            try: vol = round(float(r["Average Daily O-D Traffic (StL Volume)"]))
            except ValueError: continue
            try: tt = round(float(r["Avg Travel Time (sec)"]))
            except (KeyError, ValueError): tt = None
            flows.append([oi, di, day_types.index(dt), day_parts.index(dp), vol, tt])
    od_cfg = {
        "note": "StreetLight O-D 2087595 · Jan 2025 – Aug 2026",
        "origins": [{"name": n, "label": l, "lat": a, "lon": o} for n, l, a, o in o_zones],
        "dests":   [{"name": n, "label": l, "lat": a, "lon": o} for n, l, a, o in d_zones],
        "dayTypes": sorted(day_types), "dayParts": sorted(day_parts),
        "flows": flows,
    }
    # flows were built against first-seen order; re-map onto the sorted tables
    dt_map = {i: od_cfg["dayTypes"].index(v) for i, v in enumerate(day_types)}
    dp_map = {i: od_cfg["dayParts"].index(v) for i, v in enumerate(day_parts)}
    for fl in od_cfg["flows"]:
        fl[2], fl[3] = dt_map[fl[2]], dp_map[fl[3]]
    print(f"OD: {len(o_zones)} origins x {len(d_zones)} dests, {len(flows)} flow rows"
          + (f", UNMATCHED zone names: {missed}" if missed else ""))

# ── Crashes: statewide listings -> corridor buffer ────────────────────
crashes = filter_crashes(os.path.join(DATA, "TMC_Identification.csv"),
                         CRASH_DIR, a.buffer_ft)
years = sorted({c["year"] for c in crashes})
print(f"crashes: {len(crashes)} within {a.buffer_ft:.0f} ft ({years[0]}-{years[-1]})" if crashes else "crashes: none")

# ── Emit the corridor config ──────────────────────────────────────────
gj = json.dumps({"type": "FeatureCollection", "features": feats}, separators=(",", ":"))
cj = json.dumps(crashes, separators=(",", ":"))
sgj = json.dumps({"type": "FeatureCollection", "features": sl_feats}, separators=(",", ":"))
shj = json.dumps(sl_hourly, separators=(",", ":"))
slash2 = "/" * 2
src = (
    f"{slash2} East Arapahoe Road (CO-88), Arapahoe County CO.\n"
    f"{slash2} GENERATED by scripts/build_east_arapahoe.py - do not edit by hand.\n"
    f"{slash2} StreetLight: Corridor Study 2087472 (Jan 2025 - Aug 2026), {len(sl_feats)} directional OSM segments.\n"
    f"{slash2} INRIX: 18 TMCs, 5-min speeds -> day mean; free-flow = INRIX reference speed.\n"
    f"{slash2} Crashes: CDOT statewide listings filtered to a {a.buffer_ft:.0f} ft corridor buffer.\n"
    f"{slash2} Raw exports live in the git-ignored 'corridor data/' folder.\n"
    "window.registerCorridor({\n"
    "  id: 'east-arapahoe',\n"
    "  name: 'East Arapahoe Road',\n"
    "  agency: 'Arapahoe County, CO',\n"
    "  center: [39.5948, -104.8447],\n"
    "  zoom: 13,\n"
    "  dirs: [\n"
    "    { key:'EAST', label:'Eastbound', short:'EB', arrow:'\\u25B6' },\n"
    "    { key:'WEST', label:'Westbound', short:'WB', arrow:'\\u25C0' }\n"
    "  ],\n"
    "  streetlight: {\n"
    f"    ranges: {{ adt:[0,{adt_cap}], adtDir:{adt_cap}, speed:[20,55], tti:[1,2] }},\n"
    f"    geojson: {sgj},\n"
    f"    hourly: {shj}\n"
    "  },\n"
    f"  inrix: {{ geojson: {gj} }},\n"
    f"  tcds: {{ stations: {json.dumps(tcds_stations, separators=(',', ':'))} }},\n"
    + (f"  od: {json.dumps(od_cfg, separators=(',', ':'))},\n" if od_cfg else "")
    + f"  crashes: {cj}\n"
    "});\n"
)
with open(OUT, "w", encoding="utf-8") as f:
    f.write(src)
print(f"wrote {os.path.normpath(OUT)} ({os.path.getsize(OUT)/1024:.0f} KB)")
