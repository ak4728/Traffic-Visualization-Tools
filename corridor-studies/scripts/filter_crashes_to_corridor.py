"""Filter statewide CDOT crash listings to a corridor buffer.

Generic: works for ANY corridor defined by an INRIX TMC_Identification.csv
(each TMC's start/end coordinates form the corridor centerline segments).
Keeps every crash whose point falls within --buffer-ft of any segment,
and maps CDOT fields to the dashboard crash schema:

    { id, lat, lon, severity 0-3, sev_label, location, year, ped, bike }

Severity mapping (KABCO via the Injury 01..04 count columns):
    fatal (Number Killed or Injury 04) .. 3  "Fatal"
    suspected serious (Injury 03) ....... 3  "Serious Injury"
    suspected minor (Injury 02) ......... 2  "Injury"
    possible (Injury 01) ................ 1  "Possible Injury"
    otherwise ........................... 0  "Property Damage Only"

Usage:
    python filter_crashes_to_corridor.py --tmc "../corridor data/TMC_Identification.csv" ^
        --crash-dir "../corridor data/Statewide Crash Data" --buffer-ft 250 --out crashes.json

Raw CDOT exports are license-restricted-ish and huge — keep them in the
git-ignored "corridor data/" folder; commit only the processed output.
"""
import argparse, csv, glob, json, math, os, sys

import openpyxl


def load_segments(tmc_csv):
    """Corridor centerline as a list of ((lat1,lon1),(lat2,lon2)) segments."""
    segs = []
    with open(tmc_csv, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            try:
                segs.append(((float(row["start_latitude"]), float(row["start_longitude"])),
                             (float(row["end_latitude"]),   float(row["end_longitude"]))))
            except (KeyError, ValueError):
                continue
    if not segs:
        sys.exit("no usable segments in " + tmc_csv)
    return segs


def make_distance_fn(segs):
    """Point-to-nearest-segment distance in meters (equirectangular approx —
    fine for corridor-scale extents)."""
    lat0 = sum(s[0][0] + s[1][0] for s in segs) / (2 * len(segs))
    m_lat = 111_132.0
    m_lon = 111_320.0 * math.cos(math.radians(lat0))

    def to_xy(lat, lon):
        return lon * m_lon, lat * m_lat

    xy_segs = [(to_xy(*a), to_xy(*b)) for a, b in segs]

    def dist(lat, lon):
        px, py = to_xy(lat, lon)
        best = float("inf")
        for (ax, ay), (bx, by) in xy_segs:
            dx, dy = bx - ax, by - ay
            L2 = dx * dx + dy * dy
            t = 0.0 if L2 == 0 else max(0.0, min(1.0, ((px - ax) * dx + (py - ay) * dy) / L2))
            ex, ey = ax + t * dx - px, ay + t * dy - py
            d = math.hypot(ex, ey)
            if d < best:
                best = d
        return best

    # bounding box (degrees) for a cheap prefilter
    lats = [c for s in segs for c in (s[0][0], s[1][0])]
    lons = [c for s in segs for c in (s[0][1], s[1][1])]
    return dist, (min(lats), max(lats), min(lons), max(lons)), (m_lat, m_lon)


def classify_severity(killed, i01, i02, i03, i04):
    if killed or i04:
        return 3, "Fatal"
    if i03:
        return 3, "Serious Injury"
    if i02:
        return 2, "Injury"
    if i01:
        return 1, "Possible Injury"
    return 0, "Property Damage Only"


def _num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 0


def filter_crashes(tmc_csv, crash_dir, buffer_ft=250.0):
    segs = load_segments(tmc_csv)
    dist, (lat_lo, lat_hi, lon_lo, lon_hi), (m_lat, m_lon) = make_distance_fn(segs)
    buf_m = buffer_ft * 0.3048
    lat_pad, lon_pad = buf_m / m_lat, buf_m / m_lon

    out = []
    files = sorted(glob.glob(os.path.join(crash_dir, "*.xlsx")))
    if not files:
        sys.exit("no .xlsx files in " + crash_dir)
    for path in files:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        hdr = [str(h).strip() if h is not None else "" for h in next(rows)]
        col = {h: i for i, h in enumerate(hdr)}

        def get(r, name, default=""):
            i = col.get(name)
            return r[i] if i is not None and i < len(r) else default

        kept = 0
        for r in rows:
            try:
                lat, lon = float(get(r, "Latitude")), float(get(r, "Longitude"))
            except (TypeError, ValueError):
                continue
            if not (lat_lo - lat_pad <= lat <= lat_hi + lat_pad and
                    lon_lo - lon_pad <= lon <= lon_hi + lon_pad):
                continue
            if dist(lat, lon) > buf_m:
                continue
            sev, lab = classify_severity(_num(get(r, "Number Killed")),
                                         _num(get(r, "Injury 01")), _num(get(r, "Injury 02")),
                                         _num(get(r, "Injury 03")), _num(get(r, "Injury 04")))
            date = get(r, "Crash Date")
            year = date.year if hasattr(date, "year") else _num(str(date)[:4])
            loc1, loc2 = str(get(r, "Location 1") or "").strip(), str(get(r, "Location 2") or "").strip()
            location = (loc1 + (" & " + loc2 if loc2 else "")) or "(no location given)"
            blob = " ".join(str(get(r, c) or "") for c in
                            ("TU-1 NM Type", "TU-2 NM Type", "Crash Type")).upper()
            out.append({
                "id": str(get(r, "CUID")),
                "lat": round(lat, 6), "lon": round(lon, 6),
                "severity": sev, "sev_label": lab,
                "location": location.title(), "year": year,
                "ped": ("PED" in blob), "bike": ("BIC" in blob or "CYCL" in blob),
            })
            kept += 1
        wb.close()
        print(f"  {os.path.basename(path)}: kept {kept}")
    out.sort(key=lambda c: (c["year"], c["id"]))
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tmc", required=True, help="TMC_Identification.csv defining the corridor")
    ap.add_argument("--crash-dir", required=True, help="folder of CDOT crash-listing .xlsx files")
    ap.add_argument("--buffer-ft", type=float, default=250.0, help="buffer distance in feet (default 250)")
    ap.add_argument("--out", required=True, help="output JSON file")
    a = ap.parse_args()
    crashes = filter_crashes(a.tmc, a.crash_dir, a.buffer_ft)
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(crashes, f)
    print(f"wrote {len(crashes)} crashes within {a.buffer_ft:.0f} ft -> {a.out}")
